"""
Bộ sinh Trang bìa và Trang phụ bìa chuẩn ĐATN (Cover Generator)
Tuân thủ Phụ lục 1: Mẫu bìa Đồ án tốt nghiệp - ĐH Giao thông Vận tải
"""
import os
import re
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from .config import FONT_FAMILY

class CoverGenerator:
    @staticmethod
    def apply_cover_border(section):
        """Thêm khung viền đôi (Double Border) cho trang bìa"""
        borders = parse_xml(r'''
        <w:pgBorders %s w:offsetFrom="page">
          <w:top w:val="double" w:sz="12" w:space="24" w:color="000000"/>
          <w:left w:val="double" w:sz="12" w:space="24" w:color="000000"/>
          <w:bottom w:val="double" w:sz="12" w:space="24" w:color="000000"/>
          <w:right w:val="double" w:sz="12" w:space="24" w:color="000000"/>
        </w:pgBorders>
        ''' % nsdecls('w'))
        section._sectPr.append(borders)

    @classmethod
    def _is_compact_layout(cls, metadata: dict) -> bool:
        """Xác định xem có cần thu nhỏ kích thước để chống tràn trang hay không"""
        de_tai = metadata.get("de_tai", [])
        if isinstance(de_tai, str):
            de_tai_lines = [l.strip() for l in de_tai.splitlines() if l.strip()]
        else:
            de_tai_lines = [str(l).strip() for l in de_tai if str(l).strip()]

        thong_tin = metadata.get("thong_tin", [])
        total_rows = 0
        for item in thong_tin:
            vals = item.get("gia_tri", [])
            total_rows += max(1, len(vals) if isinstance(vals, list) else 1)

        return (total_rows >= 6) or (len(de_tai_lines) >= 2 and total_rows >= 5)

    @classmethod
    def calculate_footer_space(cls, metadata: dict) -> Pt:
        """
        Tính toán khoảng cách động (space_before) cho dòng footer '{dia_diem} – {nam}'
        sao cho dòng này được đẩy sát đáy trang (cách viền dưới ~2.0 - 2.5cm),
        tự động co dãn theo độ dài đề tài và bảng thông tin để chống tràn trang.
        """
        don_vi = metadata.get("don_vi", ["TRƯỜNG ĐẠI HỌC GIAO THÔNG VẬN TẢI", "KHOA CÔNG NGHỆ THÔNG TIN"])
        if isinstance(don_vi, str):
            don_vi = [don_vi]

        de_tai = metadata.get("de_tai", [])
        if isinstance(de_tai, str):
            de_tai_lines = [l.strip() for l in de_tai.splitlines() if l.strip()]
        else:
            de_tai_lines = [str(l).strip() for l in de_tai if str(l).strip()]
        de_tai_lines = de_tai_lines if de_tai_lines else ["TÊN ĐỀ TÀI"]

        thong_tin = metadata.get("thong_tin", [])
        total_rows = 0
        for item in thong_tin:
            vals = item.get("gia_tri", [])
            total_rows += max(1, len(vals) if isinstance(vals, list) else 1)

        is_compact = cls._is_compact_layout(metadata)

        # 1. Chiều cao header đơn vị
        h_header = len(don_vi) * 20 + 16

        # 2. Chiều cao logo (hoặc khoảng trống fallback)
        logo_path = metadata.get("logo_path", "")
        resolved_logo = logo_path
        if resolved_logo and not os.path.exists(resolved_logo):
            fallback = os.path.join(os.path.dirname(__file__), "assets", "logo.jpg")
            if os.path.exists(fallback):
                resolved_logo = fallback

        if resolved_logo and os.path.exists(resolved_logo):
            h_logo = 100 if is_compact else 135
        else:
            h_logo = 40 if is_compact else 60

        # 3. Tiêu đề loại tài liệu + Nhãn ĐỀ TÀI
        h_loai = 30 + 16
        h_nhan_detai = 18 + 6

        # 4. Chiều cao ĐỀ TÀI
        space_after_detai = 18 if is_compact else 36
        h_detai = len(de_tai_lines) * 24 + space_after_detai

        # 5. Chiều cao bảng thông tin
        row_h = 18 if is_compact else 22
        h_table = total_rows * row_h + 8

        # 6. Dòng footer
        h_footer = 18

        total_content = h_header + h_logo + h_loai + h_nhan_detai + h_detai + h_table + h_footer
        printable_height = 700  # A4 842pt - 2*71pt lề trên/dưới

        remaining = printable_height - total_content - 10
        footer_space = max(20, min(240, int(remaining)))
        return Pt(footer_space)

    @classmethod
    def render_cover_page(cls, doc: Document, metadata: dict, is_subcover: bool = False):
        """
        Sinh 1 trang bìa hoặc phụ bìa theo dữ liệu metadata linh động
        """
        don_vi = metadata.get("don_vi", ["TRƯỜNG ĐẠI HỌC GIAO THÔNG VẬN TẢI", "KHOA CÔNG NGHỆ THÔNG TIN"])
        if isinstance(don_vi, str):
            don_vi = [don_vi]

        loai = metadata.get("loai", "ĐỒ ÁN TỐT NGHIỆP").upper()
        nhan_de_tai = metadata.get("nhan_de_tai", "ĐỀ TÀI:").upper()

        de_tai = metadata.get("de_tai", [])
        if isinstance(de_tai, str):
            de_tai_lines = [l.strip() for l in de_tai.splitlines() if l.strip()]
        else:
            de_tai_lines = [str(l).strip() for l in de_tai if str(l).strip()]
        if not de_tai_lines:
            de_tai_lines = ["TÊN ĐỀ TÀI CỦA ĐỒ ÁN"]

        thong_tin = metadata.get("thong_tin", [])
        dia_diem = metadata.get("dia_diem", "Hà Nội")
        nam = metadata.get("nam", "2026")
        logo_path = metadata.get("logo_path", "assets/logo.jpg")

        is_compact = cls._is_compact_layout(metadata)

        # 1. Header cơ quan / đơn vị (hỗ trợ nhiều dòng)
        for idx, line in enumerate(don_vi):
            p_dv = doc.add_paragraph()
            p_dv.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_dv.paragraph_format.space_before = Pt(0)
            p_dv.paragraph_format.space_after = Pt(16 if idx == len(don_vi) - 1 else 3)
            r_dv = p_dv.add_run(line.upper())
            r_dv.font.name = FONT_FAMILY
            r_dv.font.size = Pt(14)
            r_dv.font.bold = True
            if is_subcover:
                r_dv.font.color.rgb = RGBColor(0, 0, 0)

        # 2. Logo trường
        resolved_logo = logo_path
        if resolved_logo and not os.path.exists(resolved_logo):
            fallback = os.path.join(os.path.dirname(__file__), "assets", "logo.jpg")
            if os.path.exists(fallback):
                resolved_logo = fallback

        if resolved_logo and os.path.exists(resolved_logo):
            p_logo = doc.add_paragraph()
            p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_logo.paragraph_format.space_before = Pt(2)
            p_logo.paragraph_format.space_after = Pt(10 if is_compact else 16)
            run_logo = p_logo.add_run()
            logo_width = Inches(1.15) if is_compact else Inches(1.5)
            run_logo.add_picture(resolved_logo, width=logo_width)
        else:
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_after = Pt(30 if is_compact else 50)

        # 3. Loại tài liệu / Báo cáo
        p_loai = doc.add_paragraph()
        p_loai.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_loai.paragraph_format.space_before = Pt(6)
        p_loai.paragraph_format.space_after = Pt(10)
        r_loai = p_loai.add_run(loai)
        r_loai.font.name = FONT_FAMILY
        r_loai.font.size = Pt(20)
        r_loai.font.bold = True
        if is_subcover:
            r_loai.font.color.rgb = RGBColor(0, 0, 0)

        # 4. Nhãn ĐỀ TÀI:
        p_nhan = doc.add_paragraph()
        p_nhan.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_nhan.paragraph_format.space_before = Pt(0)
        p_nhan.paragraph_format.space_after = Pt(4)
        r_nhan = p_nhan.add_run(nhan_de_tai)
        r_nhan.font.name = FONT_FAMILY
        r_nhan.font.size = Pt(14)
        r_nhan.font.bold = True
        if is_subcover:
            r_nhan.font.color.rgb = RGBColor(0, 0, 0)

        # 5. Nội dung ĐỀ TÀI (Đa dòng)
        space_after_detai = Pt(18 if is_compact else 36)
        for idx, dt_line in enumerate(de_tai_lines):
            p_dt = doc.add_paragraph()
            p_dt.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_dt.paragraph_format.space_before = Pt(2 if idx > 0 else 4)
            p_dt.paragraph_format.space_after = space_after_detai if idx == len(de_tai_lines) - 1 else Pt(2)
            r_dt = p_dt.add_run(dt_line.upper())
            r_dt.font.name = FONT_FAMILY
            r_dt.font.size = Pt(18)
            r_dt.font.bold = True
            if is_subcover:
                r_dt.font.color.rgb = RGBColor(0, 0, 0)
            else:
                r_dt.font.color.rgb = RGBColor(0, 32, 96)

        # 6. Chuẩn bị dữ liệu bảng thông tin (3 cột: Nhãn, Tên/Nội dung, Mã số/Phụ)
        flat_rows = []
        for item in thong_tin:
            nhan = item.get("nhan", "")
            vals = item.get("gia_tri", [])
            if not isinstance(vals, (list, tuple)):
                vals = [vals] if vals else []

            if not vals:
                flat_rows.append((nhan, "", ""))
                continue

            is_student_field = any(k in nhan.lower() for k in ["sinh viên", "sv", "thành viên", "học viên", "thực hiện"])

            for v_idx, raw_val in enumerate(vals):
                val_str = str(raw_val).strip()
                label_text = nhan if v_idx == 0 else ""

                col2_text = val_str
                col3_text = ""

                # Chỉ phân tách Tên và MSSV nếu là trường sinh viên hoặc thành viên
                if is_student_field:
                    parts = re.split(r'\s*[-–—]\s*', val_str, maxsplit=1)
                    if len(parts) == 2:
                        col2_text = parts[0].strip()
                        col3_text = parts[1].strip()

                flat_rows.append((label_text, col2_text, col3_text))

        # Dựng bảng 3 cột với viền ẩn
        if flat_rows:
            table = doc.add_table(rows=len(flat_rows), cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False

            col_widths = [Cm(5.0), Cm(6.5), Cm(3.5)]
            font_size = Pt(12 if is_compact else 13)
            pad_pt = Pt(1 if is_compact else 2)

            for r_idx, (lbl, c2, c3) in enumerate(flat_rows):
                row = table.rows[r_idx]

                # Áp dụng độ rộng cột và xóa viền ô
                for c_idx, cell in enumerate(row.cells):
                    cell.width = col_widths[c_idx]
                    tcPr = cell._tc.get_or_add_tcPr()
                    tcBorders = parse_xml(r'''
                        <w:tcBorders %s>
                            <w:top w:val="none"/>
                            <w:left w:val="none"/>
                            <w:bottom w:val="none"/>
                            <w:right w:val="none"/>
                        </w:tcBorders>
                    ''' % nsdecls('w'))
                    tcPr.append(tcBorders)

                # Gộp cột 2 và cột 3 nếu không có thông tin mã số
                if not c3:
                    row.cells[1].merge(row.cells[2])

                # Cột 1: Nhãn
                p1 = row.cells[0].paragraphs[0]
                p1.paragraph_format.space_before = pad_pt
                p1.paragraph_format.space_after = pad_pt
                p1.paragraph_format.line_spacing = 1.15
                if lbl:
                    r1 = p1.add_run(lbl)
                    r1.font.name = FONT_FAMILY
                    r1.font.size = font_size
                    r1.font.bold = True
                    if is_subcover:
                        r1.font.color.rgb = RGBColor(0, 0, 0)

                # Cột 2: Nội dung chính
                p2 = row.cells[1].paragraphs[0]
                p2.paragraph_format.space_before = pad_pt
                p2.paragraph_format.space_after = pad_pt
                p2.paragraph_format.line_spacing = 1.15
                if c2:
                    r2 = p2.add_run(c2)
                    r2.font.name = FONT_FAMILY
                    r2.font.size = font_size
                    if is_subcover:
                        r2.font.color.rgb = RGBColor(0, 0, 0)

                # Cột 3: Mã số / Phụ (nếu không bị gộp)
                if c3:
                    p3 = row.cells[2].paragraphs[0]
                    p3.paragraph_format.space_before = pad_pt
                    p3.paragraph_format.space_after = pad_pt
                    p3.paragraph_format.line_spacing = 1.15
                    r3 = p3.add_run(c3)
                    r3.font.name = FONT_FAMILY
                    r3.font.size = font_size
                    if is_subcover:
                        r3.font.color.rgb = RGBColor(0, 0, 0)

        # 7. Địa danh & Năm (Footer)
        p_footer = doc.add_paragraph()
        p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_footer.paragraph_format.space_before = cls.calculate_footer_space(metadata)
        p_footer.paragraph_format.space_after = Pt(0)
        r_footer = p_footer.add_run(f"{dia_diem} – {nam}")
        r_footer.font.name = FONT_FAMILY
        r_footer.font.size = Pt(13)
        r_footer.font.bold = True
        if is_subcover:
            r_footer.font.color.rgb = RGBColor(0, 0, 0)
